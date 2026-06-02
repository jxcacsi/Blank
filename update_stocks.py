#!/usr/bin/env python3
"""Fetch daily stock prices and update a local Excel (.xlsx) workbook.

The workbook keeps two sheets:

  * "Prices"  - an append-only daily log (one row per ticker per trading day,
                de-duplicated by Date + Ticker so re-runs never create dupes).
  * "Summary" - rebuilt on every run, showing the latest snapshot per ticker.

Data comes from the public Yahoo Finance chart API (the same source the
companion `index.html` page uses).

Usage:
    python3 update_stocks.py                      # default tickers
    python3 update_stocks.py AAPL MSFT NVDA       # explicit tickers
    TICKERS="AAPL,TSLA" python3 update_stocks.py  # via env var
    python3 update_stocks.py --output prices.xlsx --range 1mo
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_TICKERS = ["AAPL", "MSFT", "GOOGL", "TSLA", "AMZN"]
DEFAULT_OUTPUT = "stock_prices.xlsx"

PRICES_SHEET = "Prices"
SUMMARY_SHEET = "Summary"
PRICE_HEADERS = ["Date", "Ticker", "Name", "Currency",
                 "Open", "High", "Low", "Close", "Volume"]
SUMMARY_HEADERS = ["Ticker", "Name", "Last Date", "Close",
                   "Day Change", "Day Change %", "Currency"]

HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def fetch_prices(ticker: str, price_range: str = "5d") -> dict:
    """Return parsed daily price data for one ticker from Yahoo Finance."""
    url = (
        "https://query1.finance.yahoo.com/v8/finance/chart/"
        f"{urllib.parse.quote(ticker)}?range={price_range}&interval=1d"
    )
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        data = json.load(resp)

    chart = data.get("chart", {})
    if chart.get("error"):
        raise RuntimeError(chart["error"].get("description", "Unknown API error"))

    result = chart["result"][0]
    meta = result.get("meta", {})
    timestamps = result.get("timestamp") or []
    quote = result["indicators"]["quote"][0]

    rows = []
    for i, ts in enumerate(timestamps):
        close = quote["close"][i]
        if close is None:  # skip holidays / missing days
            continue
        rows.append({
            "date": dt.datetime.utcfromtimestamp(ts).date(),
            "open": quote["open"][i],
            "high": quote["high"][i],
            "low": quote["low"][i],
            "close": close,
            "volume": quote["volume"][i],
        })

    return {
        "ticker": ticker.upper(),
        "name": meta.get("shortName") or meta.get("symbol") or ticker.upper(),
        "currency": meta.get("currency") or "USD",
        "rows": rows,
    }


def load_or_create_workbook(path: str) -> Workbook:
    if os.path.exists(path):
        return load_workbook(path)
    wb = Workbook()
    # Rename the default sheet to our Prices sheet and add the header row.
    ws = wb.active
    ws.title = PRICES_SHEET
    ws.append(PRICE_HEADERS)
    return wb


def existing_keys(ws) -> set[tuple[str, str]]:
    """Return the set of (date-iso, ticker) keys already in the Prices sheet."""
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        date_val, ticker = row[0], row[1]
        date_iso = date_val.date().isoformat() if isinstance(date_val, dt.datetime) \
            else (date_val.isoformat() if isinstance(date_val, dt.date) else str(date_val))
        keys.add((date_iso, str(ticker)))
    return keys


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width + 3


def update_prices_sheet(ws, datasets: list[dict]) -> int:
    keys = existing_keys(ws)
    added = 0
    for ds in datasets:
        for r in ds["rows"]:
            key = (r["date"].isoformat(), ds["ticker"])
            if key in keys:
                continue
            ws.append([
                r["date"], ds["ticker"], ds["name"], ds["currency"],
                r["open"], r["high"], r["low"], r["close"], r["volume"],
            ])
            keys.add(key)
            added += 1

    # Sort all data rows by Date then Ticker for a tidy, stable sheet.
    data = list(ws.iter_rows(min_row=2, values_only=True))
    data = [row for row in data if row and row[0] is not None]
    data.sort(key=lambda row: (row[0], row[1]))
    ws.delete_rows(2, ws.max_row)
    for row in data:
        ws.append(row)

    # Apply formatting.
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"          # Date
        for c in row[4:8]:                            # OHLC
            c.number_format = "#,##0.00"
        row[8].number_format = "#,##0"               # Volume
    style_header(ws)
    autosize(ws)
    return added


def rebuild_summary_sheet(wb: Workbook, datasets: list[dict]) -> None:
    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    ws = wb.create_sheet(SUMMARY_SHEET, 0)  # put it first
    ws.append(SUMMARY_HEADERS)

    for ds in datasets:
        rows = ds["rows"]
        if not rows:
            continue
        last = rows[-1]
        prev_close = rows[-2]["close"] if len(rows) > 1 else None
        change = (last["close"] - prev_close) if prev_close is not None else None
        change_pct = (change / prev_close * 100) if change is not None and prev_close else None
        ws.append([
            ds["ticker"], ds["name"], last["date"], last["close"],
            change, change_pct, ds["currency"],
        ])

    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "yyyy-mm-dd"          # Last Date
        row[3].number_format = "#,##0.00"            # Close
        row[4].number_format = "#,##0.00;[Red]-#,##0.00"   # Day Change
        row[5].number_format = '0.00"%";[Red]-0.00"%"'     # Day Change %
    style_header(ws)
    autosize(ws)


def parse_args(argv: list[str]) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Update an Excel workbook with daily stock prices.")
    p.add_argument("tickers", nargs="*", help="Ticker symbols (overrides $TICKERS).")
    p.add_argument("--output", default=os.environ.get("STOCK_XLSX", DEFAULT_OUTPUT),
                   help=f"Workbook path (default: {DEFAULT_OUTPUT}).")
    p.add_argument("--range", default=os.environ.get("PRICE_RANGE", "5d"),
                   help="Yahoo range to fetch, e.g. 5d, 1mo, 1y (default: 5d).")
    return p.parse_args(argv)


def resolve_tickers(args: argparse.Namespace) -> list[str]:
    if args.tickers:
        raw = args.tickers
    elif os.environ.get("TICKERS"):
        raw = os.environ["TICKERS"].replace(",", " ").split()
    else:
        raw = DEFAULT_TICKERS
    # De-dupe while preserving order, uppercase.
    seen, out = set(), []
    for t in raw:
        t = t.strip().upper()
        if t and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    tickers = resolve_tickers(args)
    print(f"Tickers: {', '.join(tickers)}")
    print(f"Workbook: {args.output}  (range={args.range})")

    datasets, failures = [], []
    for t in tickers:
        try:
            ds = fetch_prices(t, args.range)
            if not ds["rows"]:
                raise RuntimeError("no price data returned")
            datasets.append(ds)
            last = ds["rows"][-1]
            print(f"  {t:<6} {last['date']}  close={last['close']:.2f} {ds['currency']}")
        except (urllib.error.URLError, KeyError, IndexError, RuntimeError) as e:
            failures.append(t)
            print(f"  {t:<6} FAILED: {e}", file=sys.stderr)

    if not datasets:
        print("No data fetched for any ticker; workbook left unchanged.", file=sys.stderr)
        return 1

    wb = load_or_create_workbook(args.output)
    if PRICES_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(PRICES_SHEET)
        ws.append(PRICE_HEADERS)
    added = update_prices_sheet(wb[PRICES_SHEET], datasets)
    rebuild_summary_sheet(wb, datasets)
    wb.active = wb[SUMMARY_SHEET]
    wb.save(args.output)

    print(f"Done: added {added} new price row(s); summary refreshed for "
          f"{len(datasets)} ticker(s).")
    if failures:
        print(f"Note: {len(failures)} ticker(s) failed: {', '.join(failures)}",
              file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
